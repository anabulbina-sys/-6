import tkinter as tk
from tkinter import ttk, messagebox
from shapes import rectangle, triangle, trapezoid
from docx import Document
from openpyxl import Workbook
from datetime import datetime

last_results = {}

class GeometryCalculator:
    def __init__(self, root):
        self.root = root
        root.title("Геометрический калькулятор")
        root.geometry("600x500")
        
        # Создаём вкладки
        self.notebook = ttk.Notebook(root)
        self.notebook.pack(fill='both', expand=True, padx=5, pady=5)
        
        # Вкладка Прямоугольник
        self.rect_frame = ttk.Frame(self.notebook)
        self.notebook.add(self.rect_frame, text="Прямоугольник")
        self.setup_rectangle_tab()
        
        # Вкладка Треугольник
        self.tri_frame = ttk.Frame(self.notebook)
        self.notebook.add(self.tri_frame, text="Треугольник")
        self.setup_triangle_tab()
        
        # Вкладка Трапеция
        self.trap_frame = ttk.Frame(self.notebook)
        self.notebook.add(self.trap_frame, text="Трапеция")
        self.setup_trapezoid_tab()
        
        # Кнопки сохранения внизу
        btn_frame = ttk.Frame(root)
        btn_frame.pack(fill='x', padx=5, pady=5)
        ttk.Button(btn_frame, text="Сохранить в Word", command=self.save_to_word).pack(side='left', padx=5)
        ttk.Button(btn_frame, text="Сохранить в Excel", command=self.save_to_excel).pack(side='left', padx=5)
    
    def setup_rectangle_tab(self):
        ttk.Label(self.rect_frame, text="Введите стороны прямоугольника:").pack(pady=10)
        
        frame = ttk.Frame(self.rect_frame)
        frame.pack()
        ttk.Label(frame, text="Сторона a:").grid(row=0, column=0, padx=5, pady=5)
        self.rect_a = ttk.Entry(frame)
        self.rect_a.grid(row=0, column=1, padx=5, pady=5)
        
        ttk.Label(frame, text="Сторона b:").grid(row=1, column=0, padx=5, pady=5)
        self.rect_b = ttk.Entry(frame)
        self.rect_b.grid(row=1, column=1, padx=5, pady=5)
        
        ttk.Button(self.rect_frame, text="Рассчитать", command=self.calc_rectangle).pack(pady=10)
        
        self.rect_result = ttk.Label(self.rect_frame, text="Результаты появятся здесь", justify='left')
        self.rect_result.pack(pady=10)
    
    def setup_triangle_tab(self):
        ttk.Label(self.tri_frame, text="Введите стороны треугольника:").pack(pady=10)
        
        frame = ttk.Frame(self.tri_frame)
        frame.pack()
        ttk.Label(frame, text="Сторона a:").grid(row=0, column=0, padx=5, pady=5)
        self.tri_a = ttk.Entry(frame)
        self.tri_a.grid(row=0, column=1, padx=5, pady=5)
        
        ttk.Label(frame, text="Сторона b:").grid(row=1, column=0, padx=5, pady=5)
        self.tri_b = ttk.Entry(frame)
        self.tri_b.grid(row=1, column=1, padx=5, pady=5)
        
        ttk.Label(frame, text="Сторона c:").grid(row=2, column=0, padx=5, pady=5)
        self.tri_c = ttk.Entry(frame)
        self.tri_c.grid(row=2, column=1, padx=5, pady=5)
        
        ttk.Button(self.tri_frame, text="Рассчитать", command=self.calc_triangle).pack(pady=10)
        
        self.tri_result = ttk.Label(self.tri_frame, text="Результаты появятся здесь", justify='left')
        self.tri_result.pack(pady=10)
    
    def setup_trapezoid_tab(self):
        ttk.Label(self.trap_frame, text="Введите параметры трапеции:").pack(pady=10)
        
        frame = ttk.Frame(self.trap_frame)
        frame.pack()
        ttk.Label(frame, text="Основание a:").grid(row=0, column=0, padx=5, pady=5)
        self.trap_a = ttk.Entry(frame)
        self.trap_a.grid(row=0, column=1, padx=5, pady=5)
        
        ttk.Label(frame, text="Основание b:").grid(row=1, column=0, padx=5, pady=5)
        self.trap_b = ttk.Entry(frame)
        self.trap_b.grid(row=1, column=1, padx=5, pady=5)
        
        ttk.Label(frame, text="Высота h:").grid(row=2, column=0, padx=5, pady=5)
        self.trap_h = ttk.Entry(frame)
        self.trap_h.grid(row=2, column=1, padx=5, pady=5)
        
        ttk.Button(self.trap_frame, text="Рассчитать", command=self.calc_trapezoid).pack(pady=10)
        
        self.trap_result = ttk.Label(self.trap_frame, text="Результаты появятся здесь", justify='left')
        self.trap_result.pack(pady=10)
    
    def calc_rectangle(self):
        try:
            a = float(self.rect_a.get())
            b = float(self.rect_b.get())
            
            S = rectangle.area(a, b)
            r = rectangle.inscribed_circle_radius(a, b)
            R = rectangle.circumscribed_circle_radius(a, b)
            
            text = f"Площадь: {S:.3f}\n"
            text += "Радиус вписанной окр.: " + (f"{r:.3f}" if r else "не существует") + "\n"
            text += f"Радиус описанной окр.: {R:.3f}"
            
            self.rect_result.config(text=text)
            last_results["Прямоугольник"] = {"params": f"a={a}, b={b}", "results": text.replace("\n", " | ")}
        except ValueError:
            messagebox.showerror("Ошибка", "Введите числа")
    
    def calc_triangle(self):
        try:
            a = float(self.tri_a.get())
            b = float(self.tri_b.get())
            c = float(self.tri_c.get())
            
            if a+b <= c or a+c <= b or b+c <= a:
                messagebox.showerror("Ошибка", "Треугольник не существует")
                return
            
            S = triangle.area(a, b, c)
            r = triangle.inscribed_circle_radius(a, b, c)
            R = triangle.circumscribed_circle_radius(a, b, c)
            
            text = f"Площадь: {S:.3f}\n"
            text += f"Радиус вписанной окр.: {r:.3f}\n"
            text += f"Радиус описанной окр.: {R:.3f}"
            
            self.tri_result.config(text=text)
            last_results["Треугольник"] = {"params": f"a={a}, b={b}, c={c}", "results": text.replace("\n", " | ")}
        except ValueError:
            messagebox.showerror("Ошибка", "Введите числа")
    
    def calc_trapezoid(self):
        try:
            a = float(self.trap_a.get())
            b = float(self.trap_b.get())
            h = float(self.trap_h.get())
            
            S = trapezoid.area(a, b, h)
            r = trapezoid.inscribed_circle_radius(a, b, h)
            R = trapezoid.circumscribed_circle_radius(a, b, h)
            
            text = f"Площадь: {S:.3f}\n"
            text += "Радиус вписанной окр.: " + (f"{r:.3f}" if r else "не существует") + "\n"
            text += "Радиус описанной окр.: " + (f"{R:.3f}" if R else "не существует")
            
            self.trap_result.config(text=text)
            last_results["Трапеция"] = {"params": f"a={a}, b={b}, h={h}", "results": text.replace("\n", " | ")}
        except ValueError:
            messagebox.showerror("Ошибка", "Введите числа")
    
    def save_to_word(self):
        if not last_results:
            messagebox.showinfo("Информация", "Нет данных. Выполните расчёт.")
            return
        
        doc = Document()
        doc.add_heading('Отчёт по геометрическим фигурам', 0)
        doc.add_paragraph(f'Дата: {datetime.now().strftime("%d.%m.%Y %H:%M")}')
        
        for name, data in last_results.items():
            doc.add_heading(name, level=1)
            doc.add_paragraph(f"Параметры: {data['params']}")
            for line in data['results'].split(' | '):
                doc.add_paragraph(f"• {line}")
        
        doc.save('geometry_report.docx')
        messagebox.showinfo("Успех", "Сохранено в geometry_report.docx")
    
    def save_to_excel(self):
        if not last_results:
            messagebox.showinfo("Информация", "Нет данных. Выполните расчёт.")
            return
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Результаты"
        ws['A1'], ws['B1'], ws['C1'] = "Фигура", "Параметры", "Результаты"
        
        for row, (name, data) in enumerate(last_results.items(), start=2):
            ws[f'A{row}'] = name
            ws[f'B{row}'] = data['params']
            ws[f'C{row}'] = data['results']
        
        for col in ['A', 'B', 'C']:
            ws.column_dimensions[col].width = 25
        
        wb.save('geometry_report.xlsx')
        messagebox.showinfo("Успех", "Сохранено в geometry_report.xlsx")


if __name__ == "__main__":
    root = tk.Tk()
    app = GeometryCalculator(root)
    root.mainloop()